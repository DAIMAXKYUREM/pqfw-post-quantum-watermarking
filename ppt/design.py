"""The deck's visual system.

Kept separate from the content so the two can be judged separately. The first version of
this deck used one pale-blue rounded card for every element on every slide, which is the
look people mean when they say a deck reads as generated: nothing dominates, nothing
recedes, and no choice on any slide could only have been made for *this* subject.

Three rules fix that.

**Dominance.** Deep navy carries 60-70% of the visual weight and it is dark, not tinted.
The architecture panel is a dark field with bright nodes on it, so the slide has a centre
of gravity instead of an even grey texture. Light cards are the supporting voice.

**One motif, repeated.** A hexagon badge marks every section and every stage. It is
lifted from the SIH artwork on the template itself, so it belongs to this deck rather
than being imported decoration.

**Weight, not stripes.** Emphasis comes from fill darkness, type size and colour --
never from an accent bar down the edge of a card, which is the other tell.
"""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- palette -----------------------------------------------------------------
# Navy dominates. Orange is the single sharp accent and is used sparingly enough
# that it always means "this is the mechanism". Green means a verified outcome.
NAVY_DEEP = RGBColor(0x0E, 0x1F, 0x42)
NAVY = RGBColor(0x1B, 0x33, 0x5E)
NAVY_NODE = RGBColor(0x23, 0x42, 0x76)
NAVY_EDGE = RGBColor(0x3D, 0x66, 0xAE)
INK = RGBColor(0x16, 0x1D, 0x2B)
MUTED = RGBColor(0x54, 0x62, 0x7A)
FAINT = RGBColor(0x7C, 0x8A, 0xA3)
ON_DARK = RGBColor(0xFF, 0xFF, 0xFF)
ON_DARK_DIM = RGBColor(0xA9, 0xBE, 0xDF)
ORANGE = RGBColor(0xF2, 0x79, 0x0D)
ORANGE_DEEP = RGBColor(0xC2, 0x5B, 0x02)
CYAN = RGBColor(0x1C, 0x9E, 0xB8)
GREEN = RGBColor(0x15, 0x8A, 0x5A)
GREEN_BRIGHT = RGBColor(0x22, 0xB1, 0x73)
RED = RGBColor(0xB3, 0x22, 0x2C)
AMBER = RGBColor(0xA8, 0x6A, 0x08)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CARD = RGBColor(0xF3, 0xF6, 0xFC)
CARD_EDGE = RGBColor(0xB6, 0xC8, 0xE4)
WARM = RGBColor(0xFE, 0xF4, 0xE4)
WARM_EDGE = RGBColor(0xE9, 0xC0, 0x7C)
MINT = RGBColor(0xE7, 0xF6, 0xEE)
MINT_EDGE = RGBColor(0x94, 0xCE, 0xB2)
ROSE = RGBColor(0xFD, 0xEE, 0xEE)
ROSE_EDGE = RGBColor(0xE3, 0xAF, 0xAF)

BODY = "Calibri"

# content band between the title and the blue footer bar
TOP = 1.28
BOTTOM = 6.88
LEFT = 0.40
RIGHT = 12.93


# --- primitives ---------------------------------------------------------------


def _plain(shape):
    shape.shadow.inherit = False
    if shape.has_text_frame:
        shape.text_frame.clear()
    return shape


def rect(slide, x, y, w, h, fill, edge=None, radius=0.05, line_w=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if edge is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = edge
        shape.line.width = Pt(line_w)
    return _plain(shape)


def panel(slide, x, y, w, h, fill=NAVY_DEEP, radius=0.035):
    """A dark field. The thing that stops the slide looking like a form."""
    return rect(slide, x, y, w, h, fill, edge=None, radius=radius)


def bring_to_front(shape):
    """Move a shape to the end of the shape tree.

    python-pptx appends new shapes, so anything drawn here covers the template's own
    placeholders. A dark panel silently swallowing the title placeholder is the failure
    mode; raising the placeholder afterwards is safer than lowering the panel, which
    would put it behind the template's full-bleed white rectangle instead.
    """
    element = shape._element
    parent = element.getparent()
    parent.remove(element)
    parent.append(element)
    return shape


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def para(tf, text, size=10.5, bold=False, color=INK, space_after=3, first=False,
         align=PP_ALIGN.LEFT, italic=False, spacing=None, caps=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = text.upper() if caps else text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = BODY
    if spacing is not None:
        # charSpacing is not exposed by python-pptx; set it on the run properties
        run.font._rPr.set("spc", str(int(spacing * 100)))
    return p


def rich(tf, parts, size=10.5, space_after=3, first=False, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    for text, bold, color in parts:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = BODY
    return p


def bullets(tf, items, size=10, color=INK, gap=3.5, marker="▪  ", first=True):
    """Note the ``first`` handling: without it paragraph[0] is left empty and every
    list in the deck floats a blank line below its heading."""
    for i, text in enumerate(items):
        para(tf, marker + text, size=size, color=color, space_after=gap,
             first=(first and i == 0))


# --- the motif ------------------------------------------------------------------


def hexbadge(slide, x, y, size, label, fill=ORANGE, text_color=WHITE, fsize=10):
    """The repeated mark. Lifted from the hexagon in SIH's own title artwork."""
    shape = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y),
                                   Inches(size), Inches(size * 0.88))
    shape.rotation = 90
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(fsize)
    run.font.bold = True
    run.font.color.rgb = text_color
    run.font.name = BODY
    # the shape is rotated, so the glyph has to rotate back to read upright
    shape.text_frame._txBody.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr"
    ).set("rot", "-5400000")
    return shape


def heading(slide, x, y, w, number, title, kicker=None, on_dark=False):
    """Section prompt with its badge. Every heading in the deck looks like this."""
    hexbadge(slide, x, y - 0.01, 0.30, number)
    tf = textbox(slide, x + 0.42, y - 0.03, w - 0.42, 0.34)
    para(tf, title, size=11.5, bold=True, color=ON_DARK if on_dark else NAVY,
         space_after=0, first=True)
    if kicker:
        para(tf, kicker, size=8.2, italic=True,
             color=ON_DARK_DIM if on_dark else FAINT, space_after=0)
    return tf


def label(slide, x, y, w, text, color=FAINT, size=7.4):
    tf = textbox(slide, x, y, w, 0.2)
    para(tf, text, size=size, bold=True, color=color, space_after=0, first=True,
         caps=True, spacing=1.2)
    return tf


# --- flow diagram ----------------------------------------------------------------


def node(slide, x, y, w, h, title, lines, kind="dark"):
    """One box in the architecture flow.

    Four kinds, and the difference between them carries meaning rather than variety:
    dark is a normal stage, accent is where the mechanism actually happens, good is a
    verified output, ghost is an input from outside the system.
    """
    styles = {
        "dark": (NAVY_NODE, NAVY_EDGE, ON_DARK, ON_DARK_DIM),
        "accent": (ORANGE_DEEP, ORANGE, WHITE, RGBColor(0xFF, 0xE1, 0xC2)),
        "good": (GREEN, GREEN_BRIGHT, WHITE, RGBColor(0xCC, 0xF0, 0xDF)),
        "ghost": (NAVY_DEEP, NAVY_EDGE, ON_DARK_DIM, RGBColor(0x7A, 0x92, 0xBB)),
    }
    fill, edge, tcol, scol = styles[kind]
    rect(slide, x, y, w, h, fill, edge=edge, radius=0.09, line_w=1.25)
    tf = textbox(slide, x + 0.08, y + 0.08, w - 0.16, h - 0.16, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, title, size=8.8, bold=True, color=tcol, space_after=2, first=True,
         align=PP_ALIGN.CENTER)
    for line in lines:
        para(tf, line, size=7.0, color=scol, space_after=0.5, align=PP_ALIGN.CENTER)


def flow_arrow(slide, x, y, w=0.30, color=NAVY_EDGE, thickness=2.25):
    """A real connector with a head, not a filled triangle pretending to be one."""
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                      Inches(x + w), Inches(y))
    line.line.color.rgb = color
    line.line.width = Pt(thickness)
    line.line._get_or_add_ln().append(_arrow_head())
    return line


def flow_arrow_down(slide, x, y, h=0.28, color=NAVY_EDGE, thickness=2.25):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                      Inches(x), Inches(y + h))
    line.line.color.rgb = color
    line.line.width = Pt(thickness)
    line.line._get_or_add_ln().append(_arrow_head())
    return line


def _arrow_head():
    from pptx.oxml.ns import qn
    from lxml import etree

    tail = etree.SubElement(etree.Element("dummy"), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return tail


# --- data display -----------------------------------------------------------------


def stat(slide, x, y, w, value, caption, color=NAVY, vsize=22, csize=7.4,
         on_dark=False):
    tf = textbox(slide, x, y, w, 0.78)
    para(tf, value, size=vsize, bold=True, color=color, space_after=1, first=True,
         align=PP_ALIGN.CENTER)
    para(tf, caption, size=csize, color=ON_DARK_DIM if on_dark else MUTED,
         space_after=0, align=PP_ALIGN.CENTER)


def logo(slide, path, x, y, size):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(size),
                                    Inches(size))
